from pathlib import Path

from docx import Document as DocxDocument

from app.rag.docx_audit import audit_docx_directory, audit_docx_structure


def test_audit_docx_structure_counts_body_table_header_and_footer(
    tmp_path: Path,
) -> None:
    source_dir = tmp_path / "source"
    category_dir = source_dir / "chatscm"
    category_dir.mkdir(parents=True)
    path = category_dir / "sample.docx"

    document = DocxDocument()
    document.add_paragraph("Body text")
    document.add_paragraph("")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = ""
    table.cell(1, 0).text = "B"
    table.cell(1, 1).text = "C"

    section = document.sections[0]
    section.header.paragraphs[0].text = "Header text"
    section.footer.paragraphs[0].text = "Footer text"
    document.save(path)

    stats = audit_docx_structure(path, source_dir)

    assert stats.source == "chatscm/sample.docx"
    assert stats.nonempty_body_paragraphs == 1
    assert stats.body_tables == 1
    assert stats.body_table_rows == 2
    assert stats.body_table_cells == 4
    assert stats.nonempty_body_table_cells == 3
    assert stats.sections == 1
    assert stats.nonempty_header_paragraphs == 1
    assert stats.nonempty_footer_paragraphs == 1
    assert stats.has_structural_text_outside_body_paragraphs is True


def test_audit_docx_directory_finds_docx_recursively(tmp_path: Path) -> None:
    nested = tmp_path / "a" / "b"
    nested.mkdir(parents=True)

    document = DocxDocument()
    document.add_paragraph("hello")
    document.save(nested / "one.docx")
    (nested / "ignore.txt").write_text("ignore", encoding="utf-8")

    stats = audit_docx_directory(tmp_path)

    assert [item.source for item in stats] == ["a/b/one.docx"]
