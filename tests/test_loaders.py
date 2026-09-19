from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfWriter

from app.rag.loaders import (
    discover_source_files,
    load_docx,
    load_pdf,
    load_source_documents,
)


def test_discover_source_files_is_recursive_and_filters_extensions(tmp_path: Path) -> None:
    (tmp_path / "vacinacao").mkdir()
    (tmp_path / "vacinacao" / "guia.pdf").write_bytes(b"not-a-real-pdf")
    (tmp_path / "notas.txt").write_text("ignore", encoding="utf-8")

    files = discover_source_files(tmp_path)

    assert [path.name for path in files] == ["guia.pdf"]


def test_load_docx_extracts_text_and_metadata(tmp_path: Path) -> None:
    category_dir = tmp_path / "gestacao"
    category_dir.mkdir()
    path = category_dir / "guia.docx"

    document = DocxDocument()
    document.add_paragraph("Primeiro parágrafo")
    document.add_paragraph("Segundo parágrafo")
    document.save(path)

    loaded = load_docx(path, tmp_path)

    assert len(loaded) == 1
    assert "Primeiro parágrafo" in loaded[0].page_content
    assert loaded[0].metadata["category"] == "gestacao"
    assert loaded[0].metadata["source"] == "gestacao/guia.docx"
    assert loaded[0].metadata["file_type"] == "docx"
    assert loaded[0].metadata["extraction_method"] == "docx"


def test_load_pdf_preserves_page_number_and_metadata_without_ocr(tmp_path: Path) -> None:
    category_dir = tmp_path / "vacinacao"
    category_dir.mkdir()
    path = category_dir / "calendario.pdf"

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with path.open("wb") as output:
        writer.write(output)

    loaded = load_pdf(path, tmp_path, pdf_ocr_enabled=False)

    assert len(loaded) == 1
    assert loaded[0].metadata["page"] == 1
    assert loaded[0].metadata["category"] == "vacinacao"
    assert loaded[0].metadata["source"] == "vacinacao/calendario.pdf"
    assert loaded[0].metadata["extraction_method"] == "empty"


def test_load_source_documents_keeps_processing_after_bad_file(tmp_path: Path) -> None:
    (tmp_path / "gestacao").mkdir()
    good_path = tmp_path / "gestacao" / "bom.docx"
    bad_path = tmp_path / "gestacao" / "ruim.pdf"

    document = DocxDocument()
    document.add_paragraph("conteúdo válido")
    document.save(good_path)
    bad_path.write_text("arquivo inválido", encoding="utf-8")

    report = load_source_documents(tmp_path, pdf_ocr_enabled=False)

    assert report.files_scanned == 2
    assert report.files_loaded == 1
    assert len(report.documents) == 1
    assert len(report.errors) == 1
    assert report.errors[0].source == "gestacao/ruim.pdf"
