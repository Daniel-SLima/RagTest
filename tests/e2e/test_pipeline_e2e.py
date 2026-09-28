"""Pipeline completo: ingestão → Qdrant → /v1/chat, sem providers externos.

Usa um Qdrant real quando RAGTEST_E2E_QDRANT_URL está definido (CI) e o modo local em memória
do qdrant-client nos demais casos. Embeddings e LLM são falsos e determinísticos.
"""

import hashlib
import math
import os
import re
import shutil
from pathlib import Path
from uuid import uuid4

import httpx
import pytest
import pytest_asyncio
from docx import Document as DocxDocument
from qdrant_client import AsyncQdrantClient

from app.api.dependencies import (
    get_audit_sink,
    get_conversation_service,
    get_embedding_provider,
    get_llm_provider,
    get_settings,
    get_sparse_embedding_provider,
    get_vector_store,
)
from app.conversation.service import ConversationService
from app.conversation.sqlite_store import SQLiteSessionStore
from app.core.config import Settings
from app.main import app
from app.rag.chunking import split_documents
from app.rag.embeddings.base import SparseVectorData
from app.rag.ingestion import ingest_chunks
from app.rag.loaders import load_source_documents
from app.rag.vector_store import QdrantVectorStore

pytestmark = pytest.mark.asyncio

DIM = 128
CATALOG = Path("data/source/servicos/catalogo_servicos.json")


def _tokens(text: str) -> list[str]:
    return re.findall(r"\w{3,}", text.lower())


def _bucket(token: str, size: int) -> int:
    return int(hashlib.sha1(token.encode()).hexdigest(), 16) % size


class HashEmbeddings:
    model_name = "hash-embeddings"

    async def dimension(self) -> int:
        return DIM

    def _vector(self, text: str) -> list[float]:
        vector = [0.0] * DIM
        for token in _tokens(text):
            vector[_bucket(token, DIM)] += 1.0
        norm = math.sqrt(sum(value * value for value in vector)) or 1.0
        return [value / norm for value in vector]

    async def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [self._vector(text) for text in texts]

    async def embed_query(self, text: str) -> list[float]:
        return self._vector(text)


class HashSparse:
    model_name = "hash-sparse"

    def _vector(self, text: str) -> SparseVectorData:
        counts: dict[int, float] = {}
        for token in _tokens(text):
            index = _bucket(token, 100_000)
            counts[index] = counts.get(index, 0.0) + 1.0
        return SparseVectorData(indices=list(counts), values=list(counts.values()))

    async def embed_documents(self, texts: list[str]) -> list[SparseVectorData]:
        return [self._vector(text) for text in texts]

    async def embed_query(self, text: str) -> SparseVectorData:
        return self._vector(text)


class MemorySink:
    def __init__(self) -> None:
        self.events = []

    def emit(self, event) -> None:
        self.events.append(event)


class CitingLLM:
    """Cita a primeira fonte do catálogo presente no contexto."""

    model_name = "fake-citing-llm"

    def __init__(self) -> None:
        self.calls = 0

    async def generate(self, *, system_prompt: str, user_prompt: str) -> str:
        self.calls += 1
        match = re.search(r"\[(\d+)\] Fonte: servicos/catalogo_servicos\.json", user_prompt)
        citation = match.group(1) if match else "1"
        return f"Procure a UBS mais próxima e verifique a vaga na recepção [{citation}]."


@pytest.fixture
def source_dir(tmp_path: Path) -> Path:
    root = tmp_path / "source"
    (root / "servicos").mkdir(parents=True)
    shutil.copy(CATALOG, root / "servicos" / "catalogo_servicos.json")
    (root / "chatscm").mkdir()
    faq = DocxDocument()
    faq.add_paragraph("Com quanto tempo o preventivo fica pronto?")
    faq.add_paragraph("O resultado fica pronto entre 20 e 30 dias.")
    faq.save(root / "chatscm" / "faq.docx")
    return root


@pytest_asyncio.fixture
async def qdrant_client():
    url = os.environ.get("RAGTEST_E2E_QDRANT_URL")
    client = AsyncQdrantClient(url=url) if url else AsyncQdrantClient(location=":memory:")
    try:
        yield client
    finally:
        await client.close()


@pytest_asyncio.fixture
async def indexed_store(qdrant_client, source_dir: Path) -> QdrantVectorStore:
    report = load_source_documents(source_dir)
    assert not report.errors
    chunks = split_documents(report.documents, chunk_size=1000, chunk_overlap=200)
    store = QdrantVectorStore(qdrant_client, f"e2e_{uuid4().hex[:8]}")
    stats = await ingest_chunks(
        chunks,
        embeddings=HashEmbeddings(),
        sparse_embeddings=HashSparse(),
        vector_store=store,
        upsert_batch_size=16,
        recreate=True,
    )
    assert stats.chunks_indexed == len(chunks)
    return store


@pytest_asyncio.fixture
async def client(indexed_store: QdrantVectorStore, source_dir: Path, tmp_path: Path):
    settings = Settings(
        _env_file=None,
        source_dir=source_dir,
        retrieval_auto_decompose=False,
        api_keys="seucuida:chave-e2e",
        session_db_path=tmp_path / "sessions.sqlite3",
    )
    store = SQLiteSessionStore(settings.session_db_path)
    await store.initialize()
    llm = CitingLLM()
    sink = MemorySink()
    app.dependency_overrides.update(
        {
            get_settings: lambda: settings,
            get_vector_store: lambda: indexed_store,
            get_embedding_provider: lambda: HashEmbeddings(),
            get_sparse_embedding_provider: lambda: HashSparse(),
            get_llm_provider: lambda: llm,
            get_conversation_service: lambda: ConversationService(store, settings),
            get_audit_sink: lambda: sink,
        }
    )
    app.state.audit_sink = sink
    try:
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app),
            base_url="http://e2e",
            headers={"X-API-Key": "chave-e2e"},
        ) as test_client:
            yield test_client, llm, sink
    finally:
        app.dependency_overrides.clear()
        await store.close()


async def test_chat_answers_with_citations_and_catalog_actions(client) -> None:
    test_client, llm, sink = client

    body = (
        await test_client.post(
            "/v1/chat", json={"message": "Como eu agendo a mamografia? Serviço mamografia"}
        )
    ).json()

    assert body["grounded"] is True
    assert body["display"]["status"] == "verified"
    cited = body["sources"][body["citation_ids"][0] - 1]
    assert cited["title"] == "Catálogo de serviços do Se Cuida Mulher"
    types = [action["type"] for action in body["actions"]]
    assert "open_link" in types and "schedule_reminder" in types
    assert {action["service_id"] for action in body["actions"]} == {"mamografia"}
    assert all(action["due_date"] for action in body["actions"] if action["type"] == "schedule_reminder")
    assert llm.calls == 1
    completed = [event for event in sink.events if event.event_type == "chat.completed"]
    assert len(completed) == 1 and completed[0].grounded is True


async def test_urgent_message_skips_retrieval_and_llm(client) -> None:
    test_client, llm, sink = client

    body = (
        await test_client.post("/v1/chat", json={"message": "estou grávida e sangrando"})
    ).json()

    assert body["display"]["status"] == "emergency"
    assert body["actions"][0]["type"] == "call_emergency"
    assert body["sources"] == []
    assert llm.calls == 0
    assert sink.events[-1].event_type == "chat.completed"


async def test_session_flow_persists_turns(client) -> None:
    test_client, _, _ = client

    session_id = (await test_client.post("/v1/sessions")).json()["session_id"]
    first = await test_client.post(
        "/v1/chat", json={"session_id": session_id, "message": "Quando o preventivo fica pronto?"}
    )
    history = (await test_client.get(f"/v1/sessions/{session_id}")).json()

    assert first.status_code == 200
    assert len(history["turns"]) == 1
    assert (await test_client.delete(f"/v1/sessions/{session_id}")).status_code == 204


async def test_requests_without_key_are_rejected(client) -> None:
    test_client, _, _ = client

    response = await test_client.post(
        "/v1/chat", json={"message": "oi"}, headers={"X-API-Key": "invalida"}
    )

    assert response.status_code == 401


async def test_catalog_endpoints_reflect_ingested_catalog(client) -> None:
    test_client, _, _ = client

    services = (await test_client.get("/v1/services")).json()["services"]

    assert {service["id"] for service in services} >= {"preventivo", "mamografia"}
