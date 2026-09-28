from pathlib import Path

from docx import Document as DocxDocument

from app.rag.pii_audit import PII_CATEGORIES, audit_pii_directory, find_pii


def test_find_pii_detects_brazilian_identifiers() -> None:
    text = (
        "CPF 123.456.789-09, telefone (71) 99876-5432, email maria@exemplo.com, "
        "CEP 40010-000, cartão SUS 898 0012 3456 7891 e nascida em 12/03/1990."
    )

    categories = {match.category for match in find_pii(text)}

    assert categories == {"cpf", "telefone", "email", "cep", "cns", "data"}


def test_find_pii_ignores_public_health_content() -> None:
    text = (
        "A idade recomendada para o exame preventivo é de 25 a 64 anos. "
        "A mamografia é indicada de 40 a 74 anos. Ligue 192 em caso de urgência. "
        "O 9º mês corresponde a 40 semanas e meia."
    )

    assert find_pii(text) == []


def test_audit_pii_directory_reports_counts_without_text(tmp_path: Path) -> None:
    source_dir = tmp_path / "source"
    category_dir = source_dir / "chatscm"
    category_dir.mkdir(parents=True)

    clean = DocxDocument()
    clean.add_paragraph("Procure a unidade básica de saúde mais próxima.")
    clean.save(category_dir / "limpo.docx")

    dirty = DocxDocument()
    dirty.add_paragraph("Contato: joana@exemplo.com")
    dirty.add_paragraph("CPF 123.456.789-09")
    dirty.save(category_dir / "sujo.docx")

    reports = {report.source: report for report in audit_pii_directory(source_dir)}

    assert set(reports) == {"chatscm/limpo.docx", "chatscm/sujo.docx"}
    assert reports["chatscm/limpo.docx"].total == 0
    assert reports["chatscm/sujo.docx"].counts == {"email": 1, "cpf": 1}
    assert reports["chatscm/sujo.docx"].paragraphs_with_matches == (1, 2)
    assert "joana" not in repr(reports["chatscm/sujo.docx"])


def test_pii_categories_are_stable() -> None:
    assert PII_CATEGORIES == ("cpf", "cnpj", "cns", "telefone", "email", "cep", "data")
