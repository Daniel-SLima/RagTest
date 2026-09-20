from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument


@dataclass(frozen=True, slots=True)
class DocxStructureStats:
    source: str
    body_paragraphs: int
    nonempty_body_paragraphs: int
    body_tables: int
    body_table_rows: int
    body_table_cells: int
    nonempty_body_table_cells: int
    sections: int
    header_paragraphs: int
    nonempty_header_paragraphs: int
    header_tables: int
    footer_paragraphs: int
    nonempty_footer_paragraphs: int
    footer_tables: int

    @property
    def has_structural_text_outside_body_paragraphs(self) -> bool:
        return any(
            (
                self.nonempty_body_table_cells,
                self.nonempty_header_paragraphs,
                self.header_tables,
                self.nonempty_footer_paragraphs,
                self.footer_tables,
            )
        )


def _nonempty_paragraph_count(paragraphs: object) -> int:
    return sum(
        1
        for paragraph in paragraphs
        if getattr(paragraph, "text", "").strip()
    )


def audit_docx_structure(path: Path, source_dir: Path) -> DocxStructureStats:
    document = DocxDocument(path)
    relative = path.relative_to(source_dir).as_posix()

    body_tables = list(document.tables)
    body_table_rows = sum(len(table.rows) for table in body_tables)
    body_table_cells = sum(
        len(row.cells)
        for table in body_tables
        for row in table.rows
    )
    nonempty_body_table_cells = sum(
        1
        for table in body_tables
        for row in table.rows
        for cell in row.cells
        if cell.text.strip()
    )

    header_paragraphs = 0
    nonempty_header_paragraphs = 0
    header_tables = 0
    footer_paragraphs = 0
    nonempty_footer_paragraphs = 0
    footer_tables = 0

    for section in document.sections:
        header_paragraphs += len(section.header.paragraphs)
        nonempty_header_paragraphs += _nonempty_paragraph_count(
            section.header.paragraphs
        )
        header_tables += len(section.header.tables)

        footer_paragraphs += len(section.footer.paragraphs)
        nonempty_footer_paragraphs += _nonempty_paragraph_count(
            section.footer.paragraphs
        )
        footer_tables += len(section.footer.tables)

    return DocxStructureStats(
        source=relative,
        body_paragraphs=len(document.paragraphs),
        nonempty_body_paragraphs=_nonempty_paragraph_count(document.paragraphs),
        body_tables=len(body_tables),
        body_table_rows=body_table_rows,
        body_table_cells=body_table_cells,
        nonempty_body_table_cells=nonempty_body_table_cells,
        sections=len(document.sections),
        header_paragraphs=header_paragraphs,
        nonempty_header_paragraphs=nonempty_header_paragraphs,
        header_tables=header_tables,
        footer_paragraphs=footer_paragraphs,
        nonempty_footer_paragraphs=nonempty_footer_paragraphs,
        footer_tables=footer_tables,
    )


def audit_docx_directory(source_dir: Path) -> list[DocxStructureStats]:
    source_dir = source_dir.resolve()
    paths = sorted(
        path
        for path in source_dir.rglob("*.docx")
        if path.is_file()
    )
    return [
        audit_docx_structure(path, source_dir)
        for path in paths
    ]
