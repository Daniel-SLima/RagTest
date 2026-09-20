import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf
import pytesseract
from docx import Document as DocxDocument
from langchain_core.documents import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


@dataclass(slots=True)
class LoadError:
    source: str
    message: str


@dataclass(slots=True)
class LoadReport:
    documents: list[Document] = field(default_factory=list)
    files_scanned: int = 0
    files_loaded: int = 0
    errors: list[LoadError] = field(default_factory=list)


def discover_source_files(source_dir: Path) -> list[Path]:
    if not source_dir.exists():
        return []

    return sorted(
        path
        for path in source_dir.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )


def _normalized(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    return "".join(
        character for character in normalized if not unicodedata.combining(character)
    ).lower()


def infer_audience(relative_path: Path) -> str | None:
    value = _normalized(relative_path.as_posix())

    if "pessoa_idosa/" in value or "idoso" in value:
        return "idoso"
    if "gestacao/" in value or "gestante" in value:
        return "gestante"
    if "crianca" in value:
        return "crianca"
    if "adolescent" in value or "jovens" in value:
        return "adolescente_jovem"
    if "calendario_nacional_vacinacao_adulto" in value:
        return "adulto"

    return None


def _base_metadata(path: Path, source_dir: Path) -> dict[str, str]:
    relative_path = path.relative_to(source_dir)
    category = (
        relative_path.parts[0]
        if len(relative_path.parts) > 1
        else "uncategorized"
    )

    metadata = {
        "source": relative_path.as_posix(),
        "filename": path.name,
        "category": category,
        "file_type": path.suffix.lower().lstrip("."),
    }

    audience = infer_audience(relative_path)
    if audience:
        metadata["audience"] = audience

    return metadata


def _ocr_pdf_page(
    render_document: pymupdf.Document,
    page_index: int,
    *,
    language: str,
    dpi: int,
    timeout_seconds: int,
) -> str:
    page = render_document.load_page(page_index)
    pixmap = page.get_pixmap(dpi=dpi, alpha=False)
    image = pixmap.pil_image()

    return pytesseract.image_to_string(
        image,
        lang=language,
        timeout=timeout_seconds,
    ).strip()


def load_pdf(
    path: Path,
    source_dir: Path,
    *,
    pdf_ocr_enabled: bool = False,
    pdf_ocr_language: str = "por",
    pdf_ocr_dpi: int = 200,
    pdf_ocr_timeout_seconds: int = 60,
) -> list[Document]:
    reader = PdfReader(path)
    base_metadata = _base_metadata(path, source_dir)
    extracted_texts = [(page.extract_text() or "").strip() for page in reader.pages]

    render_document: pymupdf.Document | None = None
    if pdf_ocr_enabled and any(not text for text in extracted_texts):
        render_document = pymupdf.open(path)

    documents: list[Document] = []
    try:
        for page_number, text in enumerate(extracted_texts, start=1):
            extraction_method = "text" if text else "empty"

            if not text and render_document is not None:
                text = _ocr_pdf_page(
                    render_document,
                    page_number - 1,
                    language=pdf_ocr_language,
                    dpi=pdf_ocr_dpi,
                    timeout_seconds=pdf_ocr_timeout_seconds,
                )
                extraction_method = "ocr" if text else "empty"

            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        **base_metadata,
                        "page": page_number,
                        "extraction_method": extraction_method,
                    },
                )
            )
    finally:
        if render_document is not None:
            render_document.close()

    return documents


def load_docx(path: Path, source_dir: Path) -> list[Document]:
    document = DocxDocument(path)
    text = "\n".join(
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    )

    return [
        Document(
            page_content=text,
            metadata={
                **_base_metadata(path, source_dir),
                "extraction_method": "docx",
            },
        )
    ]


def load_source_documents(
    source_dir: Path,
    *,
    pdf_ocr_enabled: bool = False,
    pdf_ocr_language: str = "por",
    pdf_ocr_dpi: int = 200,
    pdf_ocr_timeout_seconds: int = 60,
) -> LoadReport:
    source_dir = source_dir.resolve()
    files = discover_source_files(source_dir)
    report = LoadReport(files_scanned=len(files))

    for path in files:
        try:
            if path.suffix.lower() == ".pdf":
                loaded = load_pdf(
                    path,
                    source_dir,
                    pdf_ocr_enabled=pdf_ocr_enabled,
                    pdf_ocr_language=pdf_ocr_language,
                    pdf_ocr_dpi=pdf_ocr_dpi,
                    pdf_ocr_timeout_seconds=pdf_ocr_timeout_seconds,
                )
            elif path.suffix.lower() == ".docx":
                loaded = load_docx(path, source_dir)
            else:
                continue
        except Exception as exc:  # noqa: BLE001 - isolate one bad source and continue corpus loading
            report.errors.append(
                LoadError(
                    source=path.relative_to(source_dir).as_posix(),
                    message=str(exc),
                )
            )
            continue

        report.documents.extend(loaded)
        report.files_loaded += 1

    return report
