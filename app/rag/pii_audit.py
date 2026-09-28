import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as DocxDocument

PII_CATEGORIES = ("cpf", "cnpj", "cns", "telefone", "email", "cep", "data")

_PATTERNS: dict[str, re.Pattern[str]] = {
    "cpf": re.compile(r"\b\d{3}\.\d{3}\.\d{3}-\d{2}\b|\b\d{11}\b"),
    "cnpj": re.compile(r"\b\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}\b"),
    "cns": re.compile(r"\b[1-9]\d{2}[ .]?\d{4}[ .]?\d{4}[ .]?\d{4}\b"),
    "telefone": re.compile(r"(?:\(\d{2}\)\s?|\b\d{2}\s)?\b9?\d{4}-\d{4}\b"),
    "email": re.compile(r"\b[\w.+-]+@[\w-]+\.[\w.-]+\b"),
    "cep": re.compile(r"\b\d{5}-\d{3}\b"),
    "data": re.compile(r"\b\d{1,2}/\d{1,2}/\d{2,4}\b"),
}


@dataclass(frozen=True, slots=True)
class PiiMatch:
    category: str
    start: int
    end: int


@dataclass(frozen=True, slots=True)
class PiiReport:
    source: str
    paragraphs: int
    counts: dict[str, int] = field(default_factory=dict)
    paragraphs_with_matches: tuple[int, ...] = ()

    @property
    def total(self) -> int:
        return sum(self.counts.values())

    def __repr__(self) -> str:
        return (
            f"PiiReport(source={self.source!r}, paragraphs={self.paragraphs}, "
            f"counts={self.counts!r}, paragraphs_with_matches={self.paragraphs_with_matches!r})"
        )


def find_pii(text: str) -> list[PiiMatch]:
    matches: list[PiiMatch] = []
    taken: list[tuple[int, int]] = []
    for category in PII_CATEGORIES:
        for found in _PATTERNS[category].finditer(text):
            start, end = found.span()
            if any(start < t_end and end > t_start for t_start, t_end in taken):
                continue
            taken.append((start, end))
            matches.append(PiiMatch(category, start, end))
    return sorted(matches, key=lambda match: match.start)


def _docx_paragraphs(path: Path) -> list[str]:
    document = DocxDocument(path)
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    for section in document.sections:
        texts.extend(paragraph.text for paragraph in section.header.paragraphs)
        texts.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return texts


def audit_pii_docx(path: Path, source_dir: Path) -> PiiReport:
    paragraphs = _docx_paragraphs(path)
    counts: Counter[str] = Counter()
    flagged: list[int] = []
    for index, text in enumerate(paragraphs, start=1):
        found = find_pii(text)
        if found:
            flagged.append(index)
            counts.update(match.category for match in found)
    return PiiReport(
        source=path.relative_to(source_dir).as_posix(),
        paragraphs=len(paragraphs),
        counts=dict(counts),
        paragraphs_with_matches=tuple(flagged),
    )


def audit_pii_directory(source_dir: Path) -> list[PiiReport]:
    source_dir = source_dir.resolve()
    paths = sorted(path for path in source_dir.rglob("*.docx") if path.is_file())
    return [audit_pii_docx(path, source_dir) for path in paths]
